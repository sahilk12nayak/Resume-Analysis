"""Backend tests for Resume Match Analyzer API"""
import io
import os
import pytest
import requests
import docx
from pypdf import PdfWriter

BASE_URL = os.environ.get('REACT_APP_BACKEND_URL', 'https://skill-gap-analyzer-21.preview.emergentagent.com').rstrip('/')
API = f"{BASE_URL}/api"

SAMPLE_JD = """Senior Frontend Engineer — React/TypeScript
We need 5+ years experience. Required: React 18+, TypeScript, hooks, Tailwind, Redux Toolkit, REST/GraphQL,
Jest + React Testing Library, performance optimization, CI/CD with GitHub Actions. Nice to have: Next.js, design systems."""

RESUME_TEXT = """John Doe
Senior Frontend Engineer
8 years experience
Skills: React, JavaScript, TypeScript, Redux, Tailwind CSS, Jest, REST APIs
Experience: Led frontend team at Acme Corp building React 18 applications with TypeScript.
Used Redux Toolkit and React Query for state management. Wrote unit tests with Jest and RTL.
Set up CI/CD with GitHub Actions. Optimized Core Web Vitals."""


def make_docx_bytes(text: str) -> bytes:
    d = docx.Document()
    for line in text.split("\n"):
        d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_pdf_bytes(text: str) -> bytes:
    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=10)
    for line in text.split("\n"):
        pdf.cell(0, 6, line, ln=1)
    return bytes(pdf.output())


@pytest.fixture(scope="module")
def docx_bytes():
    return make_docx_bytes(RESUME_TEXT)


@pytest.fixture(scope="module")
def txt_bytes():
    return RESUME_TEXT.encode("utf-8")


# ---------- Validation tests (fast, no LLM) ----------

def test_unsupported_file_type():
    files = {"resume": ("a.png", b"fakepng", "image/png")}
    data = {"job_description": SAMPLE_JD}
    r = requests.post(f"{API}/analyze", data=data, files=files, timeout=30)
    assert r.status_code == 400, r.text
    assert "Unsupported" in r.json().get("detail", "")


def test_short_jd(txt_bytes):
    files = {"resume": ("r.txt", txt_bytes, "text/plain")}
    data = {"job_description": "too short"}
    r = requests.post(f"{API}/analyze", data=data, files=files, timeout=30)
    assert r.status_code == 400
    assert "too short" in r.json().get("detail", "").lower()


def test_empty_file():
    files = {"resume": ("r.txt", b"", "text/plain")}
    data = {"job_description": SAMPLE_JD}
    r = requests.post(f"{API}/analyze", data=data, files=files, timeout=30)
    assert r.status_code == 400
    assert "empty" in r.json().get("detail", "").lower()


def test_get_nonexistent_analysis():
    r = requests.get(f"{API}/analyses/nonexistent-id-12345", timeout=15)
    assert r.status_code == 404


# ---------- LLM-backed tests (slow) ----------

@pytest.fixture(scope="module")
def analysis_result(docx_bytes):
    files = {"resume": ("resume.docx", docx_bytes,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    data = {"job_description": SAMPLE_JD}
    r = requests.post(f"{API}/analyze", data=data, files=files, timeout=180)
    assert r.status_code == 200, f"DOCX analyze failed: {r.status_code} {r.text[:500]}"
    return r.json()


def test_analyze_docx_structure(analysis_result):
    j = analysis_result
    assert "id" in j and isinstance(j["id"], str)
    assert isinstance(j["overall_score"], int) and 0 <= j["overall_score"] <= 100
    assert "candidate" in j and "name" in j["candidate"]
    assert isinstance(j["skills"], list) and len(j["skills"]) >= 3
    assert isinstance(j["learning_plan"], list)
    s0 = j["skills"][0]
    assert s0["status"] in ("match", "partial", "missing")
    assert s0["importance"] in ("critical", "important", "nice-to-have")


def test_get_analysis_after_post(analysis_result):
    aid = analysis_result["id"]
    r = requests.get(f"{API}/analyses/{aid}", timeout=15)
    assert r.status_code == 200, r.text
    fetched = r.json()
    assert fetched["id"] == aid
    assert fetched["overall_score"] == analysis_result["overall_score"]
    # Ensure no _id leaks
    assert "_id" not in fetched


def test_analyze_txt(txt_bytes):
    files = {"resume": ("r.txt", txt_bytes, "text/plain")}
    data = {"job_description": SAMPLE_JD}
    r = requests.post(f"{API}/analyze", data=data, files=files, timeout=180)
    assert r.status_code == 200, r.text
    j = r.json()
    assert j["candidate"]["name"]


def test_analyze_pdf():
    try:
        pdf_bytes = make_pdf_bytes(RESUME_TEXT)
    except Exception as e:
        pytest.skip(f"Could not build PDF locally: {e}")
    files = {"resume": ("r.pdf", pdf_bytes, "application/pdf")}
    data = {"job_description": SAMPLE_JD}
    r = requests.post(f"{API}/analyze", data=data, files=files, timeout=180)
    assert r.status_code == 200, r.text
    j = r.json()
    assert isinstance(j["overall_score"], int)
